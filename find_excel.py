import os
from datetime import datetime, timedelta
from openpyxl import load_workbook
import pandas as pd
import docx

today = datetime.now()

# mapped drive to search.
mainpath = "G:/"
location = "Manchester"

# list of extensions to search for
exceldoc = [".xlsx", ".csv"]
worddoc = [".doc", ".docx"]

dirlist = os.listdir(mainpath)

df = pd.DataFrame(
    columns=[
        "filename",
        "extension",
        "filepath",
        "author",
        "created date",
        "changed date",
        "last accessed",
        "changed by",
        "size(KB)",
        "MatterNo",
    ]
)

# iterate over folder
for folder in dirlist:
    print("folder: " + folder)
    path = mainpath + "/" + folder
    try:
        list = os.listdir(path)
        # iterate over files in folder
        for file in list:
            # conditions
            if file.endswith(tuple(exceldoc)):
                wb = load_workbook(path + "/" + file)

                size = os.path.getsize(path + "/" + file)
                size = float(size) / 1024
                size = round(size, 2)

                extension = os.path.splitext(path + "/" + file)[1]

                lastmod = os.path.getmtime(path + "/" + file)
                lastmod = datetime.fromtimestamp(lastmod)
                lastmod = (
                    str(lastmod.day)
                    + "/"
                    + str(lastmod.month)
                    + "/"
                    + str(lastmod.year)
                )

                created = os.path.getctime(path + "/" + file)
                created = datetime.fromtimestamp(created)
                created = (
                    str(created.day)
                    + "/"
                    + str(created.month)
                    + "/"
                    + str(created.year)
                )

                access = os.path.getatime(path + "/" + file)
                access = datetime.fromtimestamp(access)
                access = (
                    str(access.day) + "/" + str(access.month) + "/" + str(access.year)
                )

                data = {
                    "filename": file,
                    "extension": extension,
                    "filepath": (path + "/" + file),
                    "author": wb.properties.creator,
                    "created date": created,
                    "changed date": lastmod,
                    "last accessed": access,
                    "changed by": wb.properties.lastModifiedBy,
                    "size(KB)": size,
                    "MatterNo":"",
                }

                df = df.append(data, ignore_index=True)
            elif file.endswith(tuple(worddoc)):
                size = os.path.getsize(path + "/" + file)
                size = float(size) / 1024
                size = round(size, 2)

                extension = os.path.splitext(path + "/" + file)[1]

                lastmod = os.path.getmtime(path + "/" + file)
                lastmod = datetime.fromtimestamp(lastmod)
                lastmod = (
                    str(lastmod.day)
                    + "/"
                    + str(lastmod.month)
                    + "/"
                    + str(lastmod.year)
                )

                created = os.path.getctime(path + "/" + file)
                created = datetime.fromtimestamp(created)
                created = (
                    str(created.day)
                    + "/"
                    + str(created.month)
                    + "/"
                    + str(created.year)
                )

                access = os.path.getatime(path + "/" + file)
                access = datetime.fromtimestamp(access)
                access = (
                    str(access.day) + "/" + str(access.month) + "/" + str(access.year)
                )

                doc = docx.Document(path + "/" + file)
                prop = doc.core_properties
                data = {
                    "filename": file,
                    "extension": extension,
                    "filepath": (path + "/" + file),
                    "author": prop.author,
                    "created date": created,
                    "changed date": lastmod,
                    "last accessed": access,
                    "changed by": prop.last_modified_by,
                    "size(KB)": size,
                    "MatterNo":"",
                }

                df = df.append(data, ignore_index=True)
            elif os.path.splitext(path + "/" + file)[1] != "":
                size = os.path.getsize(path + "/" + file)
                size = float(size) / 1024
                size = round(size, 2)
                if size > 0:
                    extension = os.path.splitext(path + "/" + file)[1]

                    lastmod = os.path.getmtime(path + "/" + file)
                    lastmod = datetime.fromtimestamp(lastmod)
                    lastmod = (
                        str(lastmod.day)
                        + "/"
                        + str(lastmod.month)
                        + "/"
                        + str(lastmod.year)
                    )

                    created = os.path.getctime(path + "/" + file)
                    created = datetime.fromtimestamp(created)
                    created = (
                        str(created.day)
                        + "/"
                        + str(created.month)
                        + "/"
                        + str(created.year)
                    )

                    access = os.path.getatime(path + "/" + file)
                    access = datetime.fromtimestamp(access)
                    access = (
                        str(access.day)
                        + "/"
                        + str(access.month)
                        + "/"
                        + str(access.year)
                    )

                    data = {
                        "filename": file,
                        "extension": extension,
                        "filepath": (path + "/" + file),
                        "author": "unknown",
                        "created date": created,
                        "changed date": lastmod,
                        "last accessed": access,
                        "changed by": "unknown",
                        "size(KB)": size,
                        "MatterNo":"",
                    }

                    df = df.append(data, ignore_index=True)
    except:
        data = {
            "filename": Exception,
            "filepath": path,
        }

print(df)
today = datetime.now()
today = str(today.day)+'-'+str(today.month)+'-'+str(today.year)
df.to_csv(location+" files docs dated "+today+".csv")